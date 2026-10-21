"""Tests for 0.11: salutation, enclosures, CV closing, DIN 5008 and update restart.
Run with:  python3 -m pytest tests/test_v011.py"""
import importlib, json, shutil, subprocess, sys
from pathlib import Path
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
if (ROOT / 'lib').is_dir():
    sys.path.insert(0, str(ROOT / 'lib'))

CASES = {
    'Frau Müller': 'Sehr geehrte Frau Müller,',
    'Herr Max Schmidt': 'Sehr geehrter Herr Schmidt,',
    'Herrn Prof. Dr. Klaus Weber': 'Sehr geehrter Herr Prof. Dr. Weber,',
    'Frau Dr. Anna von Berg': 'Sehr geehrte Frau Dr. von Berg,',
    'Frau Dr. med. Weber': 'Sehr geehrte Frau Dr. Weber,',
    'Frau Maria de la Cruz': 'Sehr geehrte Frau de la Cruz,',
    'Frau Anna Müller-Lüdenscheidt,': 'Sehr geehrte Frau Müller-Lüdenscheidt,',
    'Personalabteilung': 'Sehr geehrte Damen und Herren,',
    'Herr': 'Sehr geehrte Damen und Herren,',
    'Frau Dr.': 'Sehr geehrte Damen und Herren,',
    '': 'Sehr geehrte Damen und Herren,',
}


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setenv('RESTLOS_DATA_DIR', str(tmp_path / 'data'))
    monkeypatch.setenv('RESTLOS_UPDATE_DIR', str(tmp_path / 'updates'))
    monkeypatch.setenv('RESTLOS_LEGACY_DIR', str(tmp_path / 'legacy'))
    for name in ['storage', 'studio', 'pdf_export', 'word_export', 'letters', 'forms', 'photos', 'personal']:
        sys.modules.pop(name, None)
    return importlib.import_module('storage'), tmp_path


def sample(storage, tmp):
    from reportlab.pdfgen import canvas
    s = storage.initial(); p = storage.selected(s)
    p['profile'].update(name='Max Muster', email='max@example.org', phone='0123 456')
    p['profile']['form'].update(street='Musterweg 1', postcode='12345', city='Musterstadt')
    a = s['applications'][0]
    a.update(company='Müller Logistik GmbH', position='Lagerist', location='Musterstadt', date='21.10.2026', motivation='Ich will.', evidence='Ich kann.')
    c = canvas.Canvas(str(tmp / 'Arbeitszeugnis_Spedition_Meier.pdf')); c.drawString(100, 700, 'Zeugnis'); c.save()
    a['attachments'].append(storage.document_copy(s, tmp / 'Arbeitszeugnis_Spedition_Meier.pdf', 'Zeugnisse')['path'])
    storage.persist(s, a['id'], generate=False)
    return s, a


def pdf_text(path):
    from pypdf import PdfReader
    return '\n'.join(p.extract_text() for p in PdfReader(path).pages)


def test_salutation_rules():
    from letters import salutation_for
    for contact, expected in CASES.items():
        assert salutation_for(contact) == expected, contact


def test_salutation_python_and_javascript_match(tmp_path):
    if not shutil.which('node'):
        pytest.skip('node nicht installiert')
    script = (ROOT / 'letter-tools.js').read_text(encoding='utf-8')
    head = script.split('const letterToolsEnsure')[0]  # pure helper functions only
    runner = tmp_path / 'run.js'
    runner.write_text(head + '\nconsole.log(JSON.stringify(' + json.dumps(list(CASES)) + '.map(salutationFor)));', encoding='utf-8')
    out = json.loads(subprocess.run(['node', str(runner)], capture_output=True, text=True, check=True).stdout)
    assert out == list(CASES.values())


def test_auto_salutation_and_existing_own_salutation(env):
    storage, tmp = env
    s, a = sample(storage, tmp)
    assert a['form']['salutationMode'] == 'auto'
    a['form']['contact'] = 'Frau Dr. Anna von Berg'
    storage.persist(s, a['id'], generate=False)
    assert a['salutation'] == 'Sehr geehrte Frau Dr. von Berg,'
    # free-text recipient: first Frau/Herr line is used
    a['form']['recipientMode'] = 'text'; a['recipient'] = 'Firma\nHerrn Klaus Weber\nStraße 1'
    storage.persist(s, a['id'], generate=False)
    assert a['salutation'] == 'Sehr geehrter Herr Weber,'
    # an application with its own salutation keeps it after the update
    raw = json.loads(json.dumps(s))
    other = raw['profiles'][0]['applications'][0]
    other['salutation'] = 'Liebes Team der Müller Logistik,'
    for key in ('salutationMode', 'attachmentText'):
        other['form'].pop(key)
    loaded = storage.validate(raw)
    assert loaded['applications'][0]['form']['salutationMode'] == 'text'
    assert loaded['applications'][0]['salutation'] == 'Liebes Team der Müller Logistik,'


def test_enclosures_in_pdf_word_and_text(env):
    storage, tmp = env
    s, a = sample(storage, tmp)
    from pdf_export import make_pdf
    from word_export import make_docx
    make_pdf(s, a, tmp / 'brief.pdf', 'letter', False)
    text = pdf_text(tmp / 'brief.pdf')
    assert 'Anlagen' in text and 'Lebenslauf' in text and 'Arbeitszeugnis Spedition Meier' in text
    make_docx(s, a, tmp / 'brief.docx', 'letter')
    from docx import Document
    words = '\n'.join(p.text for p in Document(tmp / 'brief.docx').paragraphs)
    assert 'Anlagen' in words and 'Arbeitszeugnis Spedition Meier' in words
    ap = storage.application_dir(storage.selected(s), a)
    assert 'Arbeitszeugnis Spedition Meier' in (ap / 'Anschreiben' / 'Anschreiben.txt').read_text()
    a['form']['attachmentText'] = 'Zeugnisse\nZertifikate'
    make_pdf(s, a, tmp / 'brief2.pdf', 'letter', False)
    assert 'Zertifikate' in pdf_text(tmp / 'brief2.pdf') and 'Spedition' not in pdf_text(tmp / 'brief2.pdf')
    a['design']['attachmentList'] = False
    make_pdf(s, a, tmp / 'brief3.pdf', 'letter', False)
    assert 'Anlagen' not in pdf_text(tmp / 'brief3.pdf')


@pytest.mark.parametrize('template', ['Seitenleiste', 'Nordlicht', 'Klar', 'Horizont'])
def test_cv_place_date_signature(env, template):
    storage, tmp = env
    s, a = sample(storage, tmp)
    from pdf_export import make_pdf
    a['template'] = template
    make_pdf(s, a, tmp / 'cv0.pdf', 'cv', False)
    assert 'Musterstadt, 21.10.2026' not in pdf_text(tmp / 'cv0.pdf')
    a['design']['cvSignature'] = True
    make_pdf(s, a, tmp / 'cv.pdf', 'cv', False)
    assert 'Musterstadt, 21.10.2026' in pdf_text(tmp / 'cv.pdf')
    from word_export import make_docx
    make_docx(s, a, tmp / 'cv.docx', 'cv')
    from docx import Document
    assert any(p.text == 'Musterstadt, 21.10.2026' for p in Document(tmp / 'cv.docx').paragraphs)


def test_din5008_address_field_position(env):
    storage, tmp = env
    s, a = sample(storage, tmp)
    a['form'].update(contact='Frau Müller', street='Industriestraße 12', postcode='12399', city='Musterstadt')
    a['design']['letterLayout'] = 'din5008'
    storage.persist(s, a['id'], generate=False)
    from pdf_export import make_pdf
    make_pdf(s, a, tmp / 'din.pdf', 'letter', False)
    from pypdf import PdfReader
    found = {}
    def visit(text, cm, tm, font, size):
        if text.strip():
            found.setdefault(text.strip(), (tm[4] / 72 * 25.4, 297 - tm[5] / 72 * 25.4))
    PdfReader(tmp / 'din.pdf').pages[0].extract_text(visitor_text=visit)
    for line in ('Müller Logistik GmbH', 'Frau Müller', 'Industriestraße 12', '12399 Musterstadt'):
        x, y = found[line]
        assert abs(x - 25) < 0.5 and 62.7 < y < 90, (line, x, y)   # Anschriftzone, form B
    sender = next(v for k, v in found.items() if k.startswith('Max Muster ·'))
    assert 45 < sender[1] < 62.7                                  # Vermerkzone
    text = pdf_text(tmp / 'din.pdf')
    assert 'Sehr geehrte Frau Müller,' in text and 'Anlagen' in text


def test_din_check_warns_about_long_address(env):
    storage, tmp = env
    s, a = sample(storage, tmp)
    a['design']['letterLayout'] = 'din5008'; a['form']['recipientMode'] = 'text'
    a['recipient'] = '\n'.join('Zeile %d' % i for i in range(8))
    from studio import application_check
    assert any('DIN 5008' in i['message'] for i in application_check(s, a, check_size=False)['issues'])


def test_invalid_new_options_are_rejected(env):
    storage, tmp = env
    s, a = sample(storage, tmp)
    a['design']['letterLayout'] = 'quer'
    with pytest.raises(ValueError):
        storage.validate(json.loads(json.dumps(s)))
