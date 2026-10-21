import os,sys,tempfile,copy,json
from pathlib import Path
work=Path(tempfile.mkdtemp(prefix='restlos-studio-'));os.environ['RESTLOS_DATA_DIR']=str(work/'data');os.environ['RESTLOS_LEGACY_DIR']=str(work/'none');os.environ['RESTLOS_UPDATE_DIR']=str(work/'updates');sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import app,storage,studio
from pdf_export import make_pdf,TEMPLATES
from word_export import make_docx
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from docx import Document
s=storage.initial();p=s['profile'];a=s['applications'][0]
p.update(name='Mara Muster',email='mara@example.de',phone='123',address='Musterweg 1',experience='01/2020 – 12/2023\nGemeinsame Erfahrung',skills='Excel')
p['form']['experience']['mode']='text';p['form']['skills']['mode']='text';p['form']['contactMode']='text'
a.update(company='Muster GmbH',position='Assistenz',motivation='Ich organisiere Termine.',evidence='Ich bringe Erfahrung mit.')
a['form']['recipientMode']='text';a['recipient']='Muster GmbH\nMusterweg 2';a['cv']={k:p.get(k,'') for k in studio.CV_KEYS};a['cv']['experience']='01/2021 – 12/2024\nEigene Erfahrung';a['cv']['summary']='Eigenes Kurzprofil'
a2=storage.new_application();s['applications'].append(a2)
source=work/'zeugnis.pdf';c=canvas.Canvas(str(source))
for i in range(3):c.drawString(60,700,'Originalseite '+str(i+1));c.showPage()
c.save();original=source.read_bytes();d=storage.document_copy(s,source,'Zeugnisse');a['attachments']=[d['path']];a['attachmentEdits']={d['path']:[{'index':2,'rotation':90},{'index':0,'rotation':0}]}
storage.persist(s)
assert s['profile']['experience']=='01/2020 – 12/2023\nGemeinsame Erfahrung'
for name in TEMPLATES:
 a['template']=name;dest=work/(name+'.pdf');make_pdf(s,a,dest,kind='cv');text=''.join(p.extract_text() for p in PdfReader(dest).pages);assert 'Eigene Erfahrung' in text and 'Gemeinsame Erfahrung' not in text,name
make_pdf(s,a2,work/'other.pdf',kind='cv');assert 'Gemeinsame Erfahrung' in ''.join(p.extract_text() for p in PdfReader(work/'other.pdf').pages)
make_docx(s,a,work/'cv.docx');assert 'Eigene Erfahrung' in '\n'.join(p.text for p in Document(work/'cv.docx').paragraphs)
studio.attachment_pdf(a,d['path'],work/'edited.pdf');rd=PdfReader(work/'edited.pdf');assert len(rd.pages)==2 and rd.pages[0].rotation==90 and 'Originalseite 3' in rd.pages[0].extract_text();assert source.read_bytes()==original and storage.resolve(d['path']).read_bytes()==original
make_pdf(s,a,work/'all.pdf');allp=PdfReader(work/'all.pdf').pages;assert 'Originalseite 3' in allp[-2].extract_text() and 'Originalseite 1' in allp[-1].extract_text()
a['design']['breakBefore']=['experience'];a['design']['spacing']=125;a['template']='Klar';make_pdf(s,a,work/'break.pdf',kind='cv');assert len(PdfReader(work/'break.pdf').pages)>=2
make_docx(s,a,work/'break.docx');assert 'w:type="page"' in Document(work/'break.docx')._element.xml
# Missing attached originals no longer prevent saving or opening the repair dialog.
storage.resolve(d['path']).unlink();storage.persist(s,generate=False);assert len(studio.missing_documents(storage.selected(s)))==1
studio.repair_document(storage.selected(s),d['id'],source);assert not studio.missing_documents(storage.selected(s))
try:studio.repair_document(storage.selected(s),d['id'],source);raise AssertionError('overwrite')
except ValueError:pass
s['applications'][1]['company']='Andere Firma';a['motivation']='Ich hoffe das ich ich [Aufgabe] bei Andere Firma übernehme.';a['cv']['education']='12/2024 – 01/2020';storage.validate(s)
check=studio.application_check(s,a);assert any('Platzhalter' in x['message'] for x in check['issues']);assert any('Andere Firma' in x['message'] for x in check['issues']);assert any('Zeitraum' in x['message'] for x in check['issues']);assert check['pages']>=3
advice=studio.text_advice(a['letter']);assert any(x['replacement']=='ich' for x in advice);assert any('dass' in (x['replacement'] or '') for x in advice)
preview=studio.render_pages(work/'all.pdf',len(allp)-1);assert preview['start']==len(allp)-1 and len(preview['pages'])==1
storage.persist(s);loaded=storage.load();assert loaded['applications'][0]['cv']['summary']=='Eigenes Kurzprofil'
print('PASS: CV-Trennung/12 Vorlagen/Word, Seitenbearbeitung/Originalschutz, Gesamtmappe, Layout, fehlende Dateien/Reparatur, Inhaltscheck, Sprachregeln, Seitennavigation, Persistenz')
print(work)
