import os,sys,tempfile,json
from pathlib import Path
work=Path(tempfile.mkdtemp(prefix='restlos08-test-'))
os.environ['RESTLOS_DATA_DIR']=str(work/'data');os.environ['RESTLOS_UPDATE_DIR']=str(work/'updates');os.environ['RESTLOS_LEGACY_DIR']=str(work/'none')
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import app,storage,updates
from PIL import Image,ImageDraw
im=Image.new('RGB',(480,120),'white');d=ImageDraw.Draw(im);d.line([(20,90),(60,20),(70,80),(160,30),(180,85),(350,50)],fill='black',width=4);sig=work/'signature.png';im.save(sig)
updates.save_config('jurkast/restlos-bewerbung',False)
s=storage.initial();s['profiles'].append(storage.new_profile('Zweites Testprofil'));storage.persist(s)
from PySide6.QtWidgets import QMainWindow,QApplication,QFileDialog
from PySide6.QtCore import QTimer
QFileDialog.getOpenFileName=staticmethod(lambda *args:(str(sig),''))
original=QMainWindow.show;results=[]
def show(w):
 original(w);pg=w.centralWidget().page()
 QTimer.singleShot(1300,lambda:pg.runJavaScript("navigate('profile');update('p','name','Mara Muster');update('p','birthDate','1990-03-14');update('p','birthPlace','Bremen');update('p','maritalStatus','ledig');update('p','nationality','deutsch');page='editor';step='text';render();bridge('signature');"))
 def check():
  pg.runJavaScript("JSON.stringify({signature:state.profile.signature,preview:signaturePreview.length,card:!!document.querySelector('#signatureImage'),birth:state.profile.birthDate})",lambda x:results.append(json.loads(x)))
  pg.runJavaScript("(()=>{let first=state.selectedProfile;switchProfile(state.profiles[1].id);page='editor';step='text';render();let clean=!signaturePreview&&!document.querySelector('#signatureImage');switchProfile(first);page='editor';step='text';render();return JSON.stringify({clean});})()",lambda x:results.append(json.loads(x)))
  w.grab().save(str(work/'signature-ui.png'))
 QTimer.singleShot(4500,check);QTimer.singleShot(5500,QApplication.instance().quit)
QMainWindow.show=show;app.launch()
assert results and results[0]['signature'] and results[0]['preview'] and results[0]['card'],results
assert results[1]['clean'],results
s=storage.load();p=s['profile'];a=s['applications'][0]
from pdf_export import make_pdf,TEMPLATES
from word_export import make_docx
from pypdf import PdfReader
assert p['birthDate']=='1990-03-14'
for template in TEMPLATES:
 a['template']=template;path=work/(template+'.pdf');make_pdf(s,a,path,kind='cv',attachments=False);text='\n'.join(x.extract_text() for x in PdfReader(path).pages);assert '14.03.1990' in text and 'Bremen' in text and 'ledig' in text,template
for enabled in (True,False):
 a['design']['signature']=enabled;path=work/f'letter-{enabled}.pdf';make_pdf(s,a,path,kind='letter',attachments=False);pdf=PdfReader(path);assert sum(len(p.images) for p in pdf.pages)==int(enabled);assert 'Geburtsdatum' not in ''.join(p.extract_text() for p in pdf.pages)
 make_docx(s,a,work/f'letter-{enabled}.docx',kind='letter')
 from docx import Document
 doc=Document(work/f'letter-{enabled}.docx');assert len(doc.inline_shapes)==int(enabled)
make_docx(s,a,work/'cv.docx',kind='cv');assert 'Geburtsdatum: 14.03.1990' in '\n'.join(p.text for p in Document(work/'cv.docx').paragraphs)
p['birthDate']='2999-01-01'
try:storage.validate(s);raise AssertionError('Future date accepted')
except ValueError:pass
p['birthDate']='1990-03-14'
for k in ('birthDate','birthPlace','nationality','maritalStatus','signature'):p.pop(k)
storage.validate(s);assert p['birthDate']=='' and p['signature']==''
print('PASS: UI Upload/Vorschau, Persistenz, 12 CV-Vorlagen, PDF/DOCX mit und ohne Unterschrift, optionale Angaben, Altprofil-Migration, Datumsprüfung')
print('Test output:',work)
