"""Editable Word exports using a robust single-column layout."""
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from forms import materialize
from photos import cropped
from personal import details,signature_image
from studio import cv_profile
from storage import resolve
from pdf_export import LABELS
from letters import attachment_lines,wants_attachment_list,place_date,subject_line

def make_docx(s,a,path,kind='cv'):
 if kind not in ('cv','letter','cover'):raise ValueError('Für Word bitte Lebenslauf, Anschreiben oder Deckblatt wählen.')
 p=s['profile'];materialize(p,a);p=cv_profile(p,a) if kind=='cv' else p;doc=Document();section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7);section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8);section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
 normal=doc.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(a['design']['size']);normal.paragraph_format.space_after=Pt(7*a['design'].get('spacing',100)/100);normal.paragraph_format.line_spacing=1.15*a['design'].get('spacing',100)/100
 for key,size in [('Title',24),('Heading 1',12),('Heading 2',11)]:
  style=doc.styles[key];style.font.name='DejaVu Sans';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('000000');style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(7)
 doc.core_properties.author=p['name'];doc.core_properties.title={'cv':'Lebenslauf','letter':'Anschreiben','cover':'Bewerbung'}[kind]+' '+p['name'];doc.core_properties.subject=a['position'];doc.core_properties.comments='';doc.core_properties.last_modified_by=p['name']
 def add(text,style=None):
  if text.strip():return doc.add_paragraph(text,style)
 def head(with_photo=False):
  if with_photo and p.get('photo') and a['design']['photo'] and resolve(p['photo']).is_file():doc.add_picture(cropped(p),width=Cm(3.1*a['design'].get('photoScale',100)/100))
  add(p['name'] or 'Dein Name','Title')
  add(' · '.join(filter(None,[p['email'],p['phone']]))+'\n'+p['address'])
 if kind=='letter':
  head();add(a['recipient'] or a['company']);add(', '.join(filter(None,[a['location'],a['date']])));add(subject_line(a),'Heading 1');add(a['salutation'])
  for block in a['letter'].split('\n\n'):add(block)
  closing=add('Mit freundlichen Grüßen');closing.paragraph_format.keep_with_next=True
  if p.get('signature') and a['design'].get('signature',True) and resolve(p['signature']).is_file():
   stream,w,h=signature_image(p,a);par=doc.add_paragraph();par.paragraph_format.keep_with_next=True;par.add_run().add_picture(stream,width=Pt(w),height=Pt(h))
  add(p['name'])
  if wants_attachment_list(a):
   owner=next((x for x in s.get('profiles',[]) if x.get('id')==s.get('selectedProfile')),None)
   lines=attachment_lines(owner['documents'] if owner else [],a)
   if lines:
    head_par=add('Anlagen');head_par.runs[0].bold=True;head_par.paragraph_format.space_before=Pt(14);head_par.paragraph_format.keep_with_next=True
    add('\n'.join(lines))
 elif kind=='cover':
  add('Bewerbung','Title');add(a['position'],'Heading 1');add(a['company']);head(True)
 else:
  head(True);add(p['title']);add(details(p))
  for key in a['design']['order']:
   value=p.get(key,'').strip()
   if not value:continue
   if key in a['design'].get('breakBefore',[]):doc.add_page_break()
   add(LABELS[key],'Heading 1')
   for block in value.split('\n\n'):
    if key in ('experience','education'):
     # Extra blank lines in free text must not produce empty headings.
     lines=[x for x in block.strip('\n').split('\n')]
     if not lines or not lines[0].strip():continue
     par=add(lines[0]);par.runs[0].bold=True;par.paragraph_format.keep_with_next=len(lines)>1
     if len(lines)>1:add('\n'.join(lines[1:]))
    else:add(block)
  if a['design'].get('cvSignature'):
   # Place, date, signature and name at the end of the CV (0.11).
   line=add(place_date(p,a))
   if line:line.paragraph_format.space_before=Pt(18);line.paragraph_format.keep_with_next=True
   if p.get('signature') and resolve(p['signature']).is_file():
    stream,w,h=signature_image(p,a);par=doc.add_paragraph();par.paragraph_format.keep_with_next=True;par.add_run().add_picture(stream,width=Pt(w),height=Pt(h))
   add(p['name'])
 doc.save(path)
